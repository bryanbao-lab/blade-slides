#!/usr/bin/env python3
"""Audit a PPTX for BLADE V2 human-editability anti-patterns."""

import argparse
import io
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from PIL import Image

EMU_PER_INCH = 914400
GENERIC_NAME = re.compile(
    r"^(Image|Picture|Text|TextBox|Shape|Rectangle|Oval|Line|Connector|Freeform|Group|"
    r"Rounded Rectangle|Isosceles Triangle)\s*\d+$",
    re.IGNORECASE,
)
CONFUSABLE_AL = re.compile(r"\bAl\b")


def inspect_media(pptx_path):
    records = []
    with zipfile.ZipFile(pptx_path) as archive:
        for name in sorted(archive.namelist()):
            if not name.startswith("ppt/media/"):
                continue
            try:
                with Image.open(io.BytesIO(archive.read(name))) as image:
                    width, height = image.size
                    short, long = sorted((width, height))
                    records.append(
                        {
                            "name": name,
                            "width": width,
                            "height": height,
                            "longest_edge": long,
                            "low_resolution": long < 256,
                            "micro_raster_strip": short <= 8 and long >= max(80, short * 10),
                        }
                    )
            except Exception:
                records.append({"name": name, "unreadable": True})
    return records


def iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def inspect_presentation(pptx_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx_path))
    total = Counter()
    fonts = Counter()
    texts = []
    per_slide = []
    names = []
    micro_placements = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        counts = Counter()
        for shape in iter_shapes(slide.shapes):
            total["objects"] += 1
            counts["objects"] += 1
            name = shape.name or ""
            names.append(name)
            if GENERIC_NAME.match(name):
                total["generic_names"] += 1
                counts["generic_names"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                total["picture_placements"] += 1
                counts["picture_placements"] += 1
                width_in = shape.width / EMU_PER_INCH
                height_in = shape.height / EMU_PER_INCH
                if width_in < 0.18 or height_in < 0.18 or width_in * height_in < 0.035:
                    micro_placements.append(
                        {
                            "slide": slide_index,
                            "name": name,
                            "width_in": round(width_in, 4),
                            "height_in": round(height_in, 4),
                        }
                    )
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                total["native_lines"] += 1
                counts["native_lines"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                total["native_shapes"] += 1
                counts["native_shapes"] += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                total["text_bearing_shapes"] += 1
                counts["text_bearing_shapes"] += 1
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
        per_slide.append({"slide": slide_index, **dict(counts)})

    total["semantic_names"] = total["objects"] - total["generic_names"]
    return {
        "slide_count": len(prs.slides),
        "slide_size_in": [
            round(prs.slide_width / EMU_PER_INCH, 6),
            round(prs.slide_height / EMU_PER_INCH, 6),
        ],
        "totals": dict(total),
        "fonts": dict(fonts.most_common()),
        "texts": texts,
        "shape_names": names,
        "micro_picture_placements": micro_placements,
        "per_slide": per_slide,
    }


def assess(presentation, media):
    risks = []
    totals = presentation["totals"]
    objects = totals.get("objects", 0)
    pictures = totals.get("picture_placements", 0)
    generic = totals.get("generic_names", 0)
    native_lines = totals.get("native_lines", 0)
    micro_placements = len(presentation["micro_picture_placements"])
    readable_media = [item for item in media if not item.get("unreadable")]
    low_res = sum(1 for item in readable_media if item["low_resolution"])
    strips = sum(1 for item in readable_media if item["micro_raster_strip"])

    width, height = presentation["slide_size_in"]
    if height and abs(width / height - 16 / 9) < 0.03 and width > 15:
        risks.append(
            {
                "code": "oversized_canvas",
                "message": f"16:9 canvas is {width} x {height} in; verify that it is not an unnecessary doubled canvas.",
            }
        )
    if objects and generic / objects > 0.80:
        risks.append(
            {
                "code": "generic_selection_pane",
                "message": f"{generic}/{objects} objects use generic names; significant layers are not human-readable.",
            }
        )
    if pictures and micro_placements > 10 and micro_placements / pictures > 0.05:
        risks.append(
            {
                "code": "micro_picture_mosaic",
                "message": f"{micro_placements} picture placements are extremely small; likely slice-mosaic reconstruction.",
            }
        )
    if strips > 10 and native_lines < 3:
        risks.append(
            {
                "code": "rasterized_regular_lines",
                "message": f"{strips} strip-like raster assets but only {native_lines} native lines; regular geometry may be rasterized.",
            }
        )
    if readable_media and low_res / len(readable_media) > 0.40:
        risks.append(
            {
                "code": "low_resolution_fragmentation",
                "message": f"{low_res}/{len(readable_media)} media assets are below 256 px on their longest edge.",
            }
        )
    all_text = "\n".join(presentation["texts"])
    al_hits = len(CONFUSABLE_AL.findall(all_text))
    if al_hits:
        risks.append(
            {
                "code": "ai_al_confusable",
                "message": f"Found {al_hits} standalone 'Al' tokens; verify AI versus lowercase-l OCR confusion.",
            }
        )
    return risks


def main():
    parser = argparse.ArgumentParser(description="Audit BLADE PPTX human editability.")
    parser.add_argument("pptx")
    parser.add_argument("--report")
    parser.add_argument("--strict", action="store_true", help="Exit 1 when risk flags are present.")
    args = parser.parse_args()

    pptx_path = Path(args.pptx).expanduser().resolve()
    if not pptx_path.exists():
        print(json.dumps({"passed": False, "errors": [f"missing PPTX: {pptx_path}"]}, indent=2))
        sys.exit(1)

    try:
        presentation = inspect_presentation(pptx_path)
        media = inspect_media(pptx_path)
        risks = assess(presentation, media)
        readable_media = [item for item in media if not item.get("unreadable")]
        report = {
            "file": str(pptx_path),
            "passed": not risks,
            "risk_flags": risks,
            "presentation": presentation,
            "media_summary": {
                "total": len(media),
                "readable": len(readable_media),
                "low_resolution_under_256px": sum(1 for item in readable_media if item["low_resolution"]),
                "micro_raster_strips": sum(1 for item in readable_media if item["micro_raster_strip"]),
            },
            "media": media,
        }
    except Exception as exc:
        report = {"file": str(pptx_path), "passed": False, "errors": [str(exc)]}
        risks = [{"code": "audit_error", "message": str(exc)}]

    output = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(output + "\n", encoding="utf-8")
    print(output)
    sys.exit(1 if args.strict and risks else 0)


if __name__ == "__main__":
    main()
