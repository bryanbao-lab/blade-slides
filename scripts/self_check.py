#!/usr/bin/env python3
"""Self-contained smoke and negative tests for the public BLADE package."""

import json
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
INIT = ROOT / "scripts" / "init_layer_contract.py"
VALIDATE = ROOT / "scripts" / "validate_layer_contract.py"
AUDIT = ROOT / "scripts" / "audit_human_editability.py"


def run(*args, expect=0):
    result = subprocess.run(
        [sys.executable, *map(str, args)],
        check=False,
        text=True,
        capture_output=True,
    )
    if result.returncode != expect:
        raise RuntimeError(
            f"command returned {result.returncode}, expected {expect}: {args}\n"
            f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
        )
    return result


def create_assets(root):
    Image.new("RGB", (3840, 2160), (220, 230, 240)).save(root / "source.png")
    Image.new("RGB", (3840, 2160), (220, 230, 240)).save(root / "clean_base.png")

    panel = Image.new("RGBA", (900, 500), (0, 0, 0, 0))
    draw = ImageDraw.Draw(panel)
    draw.polygon(
        [
            (12, 30),
            (30, 12),
            (870, 12),
            (888, 30),
            (888, 470),
            (870, 488),
            (30, 488),
            (12, 470),
        ],
        fill=(255, 252, 244, 235),
        outline=(202, 143, 35, 255),
        width=6,
    )
    panel.save(root / "panel_shell.png")

    cropped = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
    bad_draw = ImageDraw.Draw(cropped)
    bad_draw.rectangle(
        (0, 0, 199, 99),
        fill=(255, 255, 255, 230),
        outline=(210, 140, 30, 255),
        width=5,
    )
    cropped.save(root / "bad_cropped_panel.png")


def create_pptx(root):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    base = slide.shapes.add_picture(
        str(root / "clean_base.png"), 0, 0, prs.slide_width, prs.slide_height
    )
    base.name = "s99_clean_base"

    panel = slide.shapes.add_picture(
        str(root / "panel_shell.png"),
        Inches(1.0),
        Inches(1.7),
        Inches(4.5),
        Inches(2.5),
    )
    panel.name = "s99_panel_shell"

    title = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.35), Inches(6), Inches(0.7)
    )
    title.name = "s99_title"
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "Verified title"
    paragraph.runs[0].font.size = Pt(30)

    prs.save(root / "page.pptx")


def complete_contract(root):
    contract_path = root / "layer_contract.json"
    run(
        INIT,
        "--slide-id",
        "slide_99",
        "--source",
        root / "source.png",
        "--out",
        contract_path,
        "--mode",
        "full_rebuild",
    )
    data = json.loads(contract_path.read_text(encoding="utf-8"))
    data["structure_policy"]["selection_pane_reviewed"] = True
    data["content_verification"]["source_zoom_reviewed"] = True
    data["content_verification"]["proper_nouns"] = ["Verified title"]
    data["layers"].extend(
        [
            {
                "id": "panel_shell",
                "kind": "transparent_png",
                "asset_role": "panel_shell",
                "path": str(root / "panel_shell.png"),
                "bbox_px": [288, 490, 1296, 720],
                "z_index": 20,
                "editable": True,
                "significant": True,
                "powerpoint_name": "s99_panel_shell",
                "movement_unit": "complete panel shell",
                "reuse_key": "panel_shell_primary",
                "closed_border_required": True,
                "safety_padding_px": 4,
                "semantic_group": "panels",
            },
            {
                "id": "title",
                "kind": "native_text",
                "text": "Verified title",
                "bbox_px": [144, 101, 1728, 202],
                "z_index": 50,
                "editable": True,
                "significant": True,
                "powerpoint_name": "s99_title",
                "movement_unit": "slide title",
                "semantic_group": "title",
            },
        ]
    )
    data["expected"].update(
        {
            "native_text_min": 1,
            "picture_min": 2,
            "movable_foreground_min": 1,
            "semantic_named_significant_min": 3,
        }
    )
    contract_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return data, contract_path


def main():
    skill_text = (ROOT / "SKILL.md").read_text(encoding="utf-8")
    if not skill_text.startswith("---\n") or "name: blade-slides" not in skill_text:
        raise RuntimeError("SKILL.md front matter is missing or invalid")

    with tempfile.TemporaryDirectory(prefix="blade-slides-self-check-") as temp:
        root = Path(temp)
        create_assets(root)
        create_pptx(root)
        data, contract_path = complete_contract(root)

        run(
            VALIDATE,
            contract_path,
            "--pptx",
            root / "page.pptx",
            "--report",
            root / "validation.json",
        )
        run(
            AUDIT,
            root / "page.pptx",
            "--strict",
            "--report",
            root / "human_editability.json",
        )

        bad = json.loads(json.dumps(data))
        for layer in bad["layers"]:
            if layer["id"] == "panel_shell":
                layer["path"] = str(root / "bad_cropped_panel.png")
        bad_path = root / "bad_contract.json"
        bad_path.write_text(
            json.dumps(bad, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        negative = run(VALIDATE, bad_path, expect=1)
        if "transparent canvas corners contain visible pixels" not in negative.stdout:
            raise RuntimeError("negative test did not catch visible corner contamination")
        if "visible edge padding 0px is below 4px" not in negative.stdout:
            raise RuntimeError("negative test did not catch cropped border padding")

    print("BLADE Slides public package self-check: PASS")


if __name__ == "__main__":
    main()
