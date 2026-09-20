#!/usr/bin/env python3
import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from PIL import Image

ALLOWED_KINDS = {
    "clean_base",
    "native_text",
    "native_shape",
    "transparent_png",
    "fixed_scene_visual",
}
ALLOWED_MODES = {"full_rebuild", "local_refinement", "review_draft"}
GENERIC_NAME = re.compile(
    r"^(Image|Picture|Text|TextBox|Shape|Rectangle|Oval|Line|Connector|Freeform|Group|"
    r"Rounded Rectangle|Isosceles Triangle)\s*\d+$",
    re.IGNORECASE,
)


def resolve_path(value, contract_path):
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = contract_path.parent / path
    return path.resolve()


def valid_string_list(value):
    return isinstance(value, list) and all(isinstance(item, str) and item for item in value)


def validate_contract(contract_path):
    errors, warnings = [], []
    data = json.loads(contract_path.read_text(encoding="utf-8"))
    schema_version = data.get("schema_version", 1)
    size_px = data.get("slide_size_px")
    if not isinstance(size_px, list) or len(size_px) != 2:
        width, height = 0, 0
        errors.append("slide_size_px must contain positive width and height")
    else:
        width, height = size_px
        if not isinstance(width, (int, float)) or not isinstance(height, (int, float)) or width <= 0 or height <= 0:
            errors.append("slide_size_px must contain positive width and height")

    source_image = data.get("source_image")
    if not source_image:
        errors.append("source_image is required")
    elif not resolve_path(source_image, contract_path).exists():
        errors.append(f"missing source image: {resolve_path(source_image, contract_path)}")

    if schema_version >= 3:
        if data.get("skill_version") != "2.0.0":
            errors.append("schema_version 3 requires skill_version 2.0.0")
        mode = data.get("reconstruction_mode")
        if mode not in ALLOWED_MODES:
            errors.append(f"unsupported reconstruction_mode: {mode}")
        delivery = data.get("delivery_level")
        if mode == "review_draft" and delivery != "review_draft":
            errors.append("review_draft mode must use delivery_level review_draft")
        if mode != "review_draft" and delivery != "blade_full":
            errors.append("final BLADE modes must use delivery_level blade_full")
        size_in = data.get("slide_size_in")
        if (
            not isinstance(size_in, list)
            or len(size_in) != 2
            or any(not isinstance(value, (int, float)) for value in size_in)
            or min(size_in) <= 0
        ):
            errors.append("slide_size_in must contain positive width and height")

        baseline = data.get("accepted_baseline")
        if not isinstance(baseline, dict):
            errors.append("accepted_baseline is required")
        elif mode == "local_refinement":
            if not baseline.get("source_deck"):
                errors.append("local_refinement requires accepted_baseline.source_deck")
            if not baseline.get("allowed_change_slides"):
                errors.append("local_refinement requires allowed_change_slides")
            if not baseline.get("protected_slides"):
                errors.append("local_refinement requires protected_slides")
            if baseline.get("protected_render_compare_required") is not True:
                errors.append("local_refinement requires protected render comparison")

        structure = data.get("structure_policy")
        if not isinstance(structure, dict):
            errors.append("structure_policy is required")
        else:
            for key in (
                "human_editable_layer_model",
                "semantic_names_required",
                "regular_geometry_native",
                "micro_raster_fragments_forbidden",
                "meaningful_module_granularity",
                "selection_pane_reviewed",
            ):
                if structure.get(key) is not True:
                    errors.append(f"structure_policy.{key} must be true")

    editability = data.get("editability_policy")
    if not isinstance(editability, dict):
        errors.append("editability_policy is required")
    else:
        for key in ("user_requested_movable", "preserve_source_pixels", "hero_visuals"):
            if not valid_string_list(editability.get(key)):
                errors.append(f"editability_policy.{key} must be a list of non-empty strings")
        if editability.get("script_generated_visuals_forbidden") is not True:
            errors.append("editability_policy.script_generated_visuals_forbidden must be true")

    fidelity = data.get("visual_fidelity_checks")
    required_fidelity = [
        "nested_corner_pairs",
        "optical_material_regions",
        "open_gradient_regions",
        "repeated_system_motifs",
        "icon_semantic_anchors",
        "copy_linebreak_locks",
        "high_design_paths",
        "perspective_bound_scene_elements",
        "occlusion_relationships",
        "hero_composite_boundaries",
        "layer_exclusivity_pairs",
        "collision_clearances",
    ]
    if schema_version >= 3:
        required_fidelity += [
            "edge_completeness_checks",
            "contained_shape_insets",
            "paired_panel_seams",
            "repeated_grid_systems",
            "asset_safety_padding",
        ]
    if not isinstance(fidelity, dict):
        errors.append("visual_fidelity_checks is required")
    else:
        for key in required_fidelity:
            if not valid_string_list(fidelity.get(key)):
                errors.append(f"visual_fidelity_checks.{key} must be a list of non-empty strings")

    content = data.get("content_verification")
    if not isinstance(content, dict):
        errors.append("content_verification is required")
    else:
        if content.get("source_zoom_reviewed") is not True:
            errors.append("content_verification.source_zoom_reviewed must be true")
        proper_nouns = content.get("proper_nouns")
        if not valid_string_list(proper_nouns) or not proper_nouns:
            errors.append("content_verification.proper_nouns must contain verified source strings")
        if schema_version >= 3:
            confusables = content.get("confusable_pairs_checked")
            if not valid_string_list(confusables) or len(confusables) < 2:
                errors.append("content_verification.confusable_pairs_checked needs at least two checks")
            forbidden = content.get("forbidden_text_tokens")
            if not valid_string_list(forbidden):
                errors.append("content_verification.forbidden_text_tokens must be a string list")

    layers = data.get("layers", [])
    if not layers:
        errors.append("layers is empty")
    ids = [layer.get("id") for layer in layers]
    if any(not value for value in ids):
        errors.append("every layer requires a non-empty id")
    duplicates = [item for item, count in Counter(ids).items() if count > 1]
    if duplicates:
        errors.append(f"duplicate layer ids: {duplicates}")

    source_critical_ids = [layer.get("id") for layer in layers if layer.get("source_critical") is True]
    if source_critical_ids:
        source_fidelity = data.get("source_fidelity")
        if not isinstance(source_fidelity, dict):
            errors.append("source_fidelity is required when source_critical layers exist")
        else:
            if source_fidelity.get("source_critical_reviewed") is not True:
                errors.append("source_fidelity.source_critical_reviewed must be true")
            evidence = source_fidelity.get("review_evidence")
            if not valid_string_list(evidence) or not evidence:
                errors.append("source_fidelity.review_evidence must list evidence files")
            else:
                for value in evidence:
                    if not resolve_path(value, contract_path).exists():
                        errors.append(f"missing source-critical review evidence: {value}")
            specs = source_fidelity.get("layers")
            if not isinstance(specs, dict):
                errors.append("source_fidelity.layers must be an object")
            else:
                for layer_id in source_critical_ids:
                    spec = specs.get(layer_id)
                    if not isinstance(spec, dict):
                        errors.append(f"{layer_id}: missing source_fidelity.layers entry")
                        continue
                    if spec.get("must_not_derive_from_clean_base") is not True:
                        errors.append(f"{layer_id}: must_not_derive_from_clean_base must be true")
                    anchors = spec.get("visual_anchors")
                    if not valid_string_list(anchors) or len(anchors) < 3:
                        errors.append(f"{layer_id}: visual_anchors must contain at least three checks")
                    has_annotation = spec.get("has_semantic_annotation")
                    if has_annotation is not None and not isinstance(has_annotation, bool):
                        errors.append(f"{layer_id}: has_semantic_annotation must be true or false")
                    if has_annotation is True:
                        targets = spec.get("annotation_targets")
                        if not isinstance(targets, list) or not targets:
                            errors.append(f"{layer_id}: annotation_targets are required")
                        else:
                            for index, target in enumerate(targets):
                                prefix = f"{layer_id}: annotation_targets[{index}]"
                                if not isinstance(target, dict):
                                    errors.append(f"{prefix} must be an object")
                                    continue
                                if not target.get("label") or not target.get("target"):
                                    errors.append(f"{prefix} requires label and target")
                                for bbox_key in ("source_bbox_px", "render_bbox_px"):
                                    bbox = target.get(bbox_key)
                                    if (
                                        not isinstance(bbox, list)
                                        or len(bbox) != 4
                                        or any(not isinstance(value, (int, float)) for value in bbox)
                                    ):
                                        errors.append(f"{prefix}.{bbox_key} must contain four numbers")
                                ratio = target.get("min_visible_ratio")
                                if not isinstance(ratio, (int, float)) or not 0.8 <= ratio <= 1:
                                    errors.append(f"{prefix}.min_visible_ratio must be between 0.8 and 1.0")
                                if target.get("connection_verified") is not True:
                                    errors.append(f"{prefix}.connection_verified must be true")

    layer_counts = Counter()
    significant_names = []
    for layer in layers:
        layer_id = layer.get("id", "<unknown>")
        kind = layer.get("kind")
        layer_counts[kind] += 1
        if kind not in ALLOWED_KINDS:
            errors.append(f"{layer_id}: unsupported kind {kind}")
        bbox = layer.get("bbox_px")
        if not isinstance(bbox, list) or len(bbox) != 4:
            errors.append(f"{layer_id}: bbox_px must have four numbers")
        else:
            x, y, w, h = bbox
            if any(not isinstance(v, (int, float)) for v in bbox) or min(x, y, w, h) < 0 or w <= 0 or h <= 0:
                errors.append(f"{layer_id}: invalid bbox {bbox}")
            elif width and height and (x + w > width or y + h > height):
                errors.append(f"{layer_id}: bbox exceeds slide bounds")

        if schema_version >= 3 and layer.get("significant") is True:
            name = layer.get("powerpoint_name")
            movement = layer.get("movement_unit")
            if not isinstance(name, str) or not name or GENERIC_NAME.match(name):
                errors.append(f"{layer_id}: significant layer needs a semantic powerpoint_name")
            else:
                significant_names.append(name)
            if not isinstance(movement, str) or not movement:
                errors.append(f"{layer_id}: significant layer needs movement_unit")

        if kind in {"clean_base", "transparent_png", "fixed_scene_visual"}:
            value = layer.get("path")
            if not value:
                errors.append(f"{layer_id}: image layer requires path")
                continue
            path = resolve_path(value, contract_path)
            if not path.exists():
                errors.append(f"{layer_id}: missing image {path}")
                continue
            try:
                with Image.open(path) as image:
                    if kind == "clean_base" and image.size != (width, height):
                        errors.append(f"{layer_id}: clean base is {image.size}, expected {(width, height)}")
                    if kind == "transparent_png":
                        if image.format != "PNG" or "A" not in image.getbands():
                            errors.append(f"{layer_id}: transparent asset must be an RGBA PNG")
                            continue
                        alpha = image.getchannel("A")
                        lo, hi = alpha.getextrema()
                        if lo == 255:
                            errors.append(f"{layer_id}: alpha channel is fully opaque")
                        if hi == 0:
                            errors.append(f"{layer_id}: asset is fully transparent")
                        corner_alpha = max(
                            alpha.crop((0, 0, min(3, image.width), min(3, image.height))).getextrema()[1],
                            alpha.crop((max(0, image.width - 3), 0, image.width, min(3, image.height))).getextrema()[1],
                            alpha.crop((0, max(0, image.height - 3), min(3, image.width), image.height)).getextrema()[1],
                            alpha.crop((max(0, image.width - 3), max(0, image.height - 3), image.width, image.height)).getextrema()[1],
                        )
                        if corner_alpha > 0 and layer.get("allow_opaque_corners") is not True:
                            errors.append(f"{layer_id}: transparent canvas corners contain visible pixels")
                        visible_pixels = sum(alpha.histogram()[16:])
                        visible_fraction = visible_pixels / (image.width * image.height)
                        is_line_icon = layer.get("asset_role") == "line_icon" or str(layer_id).startswith("icon_")
                        if is_line_icon and visible_fraction > 0.55 and layer.get("allow_dense_alpha") is not True:
                            errors.append(f"{layer_id}: visible alpha occupancy {visible_fraction:.3f} suggests matte contamination")
                        visible_bbox = alpha.point(lambda v: 255 if v >= 32 else 0).getbbox()
                        if is_line_icon and visible_bbox:
                            span = max(
                                (visible_bbox[2] - visible_bbox[0]) / image.width,
                                (visible_bbox[3] - visible_bbox[1]) / image.height,
                            )
                            if span < 0.65 and layer.get("allow_small_alpha") is not True:
                                errors.append(f"{layer_id}: visible alpha span {span:.3f} suggests excessive padding")
                        if layer.get("closed_border_required") is True and visible_bbox:
                            required = max(2, int(layer.get("safety_padding_px", 2)))
                            padding = min(
                                visible_bbox[0],
                                visible_bbox[1],
                                image.width - visible_bbox[2],
                                image.height - visible_bbox[3],
                            )
                            if padding < required:
                                errors.append(f"{layer_id}: visible edge padding {padding}px is below {required}px")
                        short, long = sorted(image.size)
                        if short <= 8 and long >= short * 10 and layer.get("asset_role") != "line_texture":
                            errors.append(f"{layer_id}: micro-raster strip is forbidden; use native geometry")
            except Exception as exc:
                errors.append(f"{layer_id}: cannot inspect image: {exc}")

    duplicates = [name for name, count in Counter(significant_names).items() if count > 1]
    if duplicates:
        errors.append(f"duplicate significant powerpoint_name values: {duplicates}")
    movable_foreground = sum(
        1
        for layer in layers
        if layer.get("editable") is True and layer.get("kind") not in {"clean_base", "fixed_scene_visual"}
    )
    movable_min = data.get("expected", {}).get("movable_foreground_min")
    if movable_min is not None and movable_foreground < movable_min:
        errors.append(f"movable foreground layers={movable_foreground}, expected at least {movable_min}")
    layer_counts["movable_foreground"] = movable_foreground
    return data, layer_counts, errors, warnings


def inspect_pptx(pptx_path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(pptx_path))
    counts = Counter()
    texts, names = [], []
    for slide in prs.slides:
        for shape in slide.shapes:
            names.append(shape.name)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["picture"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                counts["native_line"] += 1
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                counts["native_shape"] += 1
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                counts["native_text"] += 1
                texts.append(shape.text)
    counts["slide_count"] = len(prs.slides)
    counts["generic_names"] = sum(1 for name in names if GENERIC_NAME.match(name or ""))
    counts["semantic_names"] = len(names) - counts["generic_names"]
    return prs, counts, texts, names


def normalize_text(value):
    return " ".join(str(value).split())


def main():
    parser = argparse.ArgumentParser(description="Validate a BLADE Slides layer contract and optional PPTX.")
    parser.add_argument("contract")
    parser.add_argument("--pptx")
    parser.add_argument("--report")
    args = parser.parse_args()

    contract_path = Path(args.contract).expanduser().resolve()
    data, layer_counts, errors, warnings = validate_contract(contract_path)
    pptx_counts = None
    if args.pptx:
        pptx_path = Path(args.pptx).expanduser().resolve()
        if not pptx_path.exists():
            errors.append(f"missing PPTX: {pptx_path}")
        else:
            try:
                prs, pptx_counts, pptx_texts, pptx_names = inspect_pptx(pptx_path)
                expected = data.get("expected", {})
                exact = expected.get("slide_count")
                if exact is not None and pptx_counts["slide_count"] != exact:
                    errors.append(f"PPTX slide_count={pptx_counts['slide_count']}, expected {exact}")
                for count_key, expected_key in (
                    ("native_text", "native_text_min"),
                    ("native_shape", "native_shape_min"),
                    ("native_line", "native_line_min"),
                    ("picture", "picture_min"),
                ):
                    minimum = expected.get(expected_key)
                    if minimum is not None and pptx_counts[count_key] < minimum:
                        errors.append(f"PPTX {count_key}={pptx_counts[count_key]}, expected at least {minimum}")
                size_in = data.get("slide_size_in")
                if isinstance(size_in, list) and len(size_in) == 2:
                    actual = [prs.slide_width / 914400, prs.slide_height / 914400]
                    if any(abs(a - b) > 0.01 for a, b in zip(actual, size_in)):
                        errors.append(f"PPTX slide size {actual} does not match contract {size_in}")
                normalized_shapes = [normalize_text(v) for v in pptx_texts]
                normalized_all = "\n".join(normalized_shapes)
                for layer in data.get("layers", []):
                    if layer.get("kind") == "native_text" and layer.get("text"):
                        expected_text = normalize_text(layer["text"])
                        if expected_text not in normalized_shapes:
                            errors.append(f"PPTX missing native_text layer {layer.get('id')}: {expected_text}")
                    if layer.get("significant") is True and layer.get("powerpoint_name") not in pptx_names:
                        errors.append(f"PPTX missing semantic shape name: {layer.get('powerpoint_name')}")
                for value in data.get("content_verification", {}).get("proper_nouns", []):
                    if normalize_text(value) not in normalized_all:
                        errors.append(f"PPTX missing verified source string: {value}")
                for token in data.get("content_verification", {}).get("forbidden_text_tokens", []):
                    if token and token in normalized_all:
                        errors.append(f"PPTX contains forbidden text token: {token}")
                minimum_names = expected.get("semantic_named_significant_min")
                if minimum_names is not None:
                    present = sum(
                        1
                        for layer in data.get("layers", [])
                        if layer.get("significant") is True and layer.get("powerpoint_name") in pptx_names
                    )
                    if present < minimum_names:
                        errors.append(f"PPTX semantic significant names={present}, expected at least {minimum_names}")
            except Exception as exc:
                errors.append(f"cannot inspect PPTX: {exc}")

    report = {
        "contract": str(contract_path),
        "schema_version": data.get("schema_version"),
        "passed": not errors,
        "layer_counts": dict(layer_counts),
        "pptx_counts": dict(pptx_counts) if pptx_counts else None,
        "errors": errors,
        "warnings": warnings,
    }
    output = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(output + "\n", encoding="utf-8")
    print(output)
    sys.exit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
